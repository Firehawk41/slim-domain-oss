"""
Parser tests use in-memory openpyxl workbooks and a stub resolver so they
never need live domain services or a real Excel file.
"""
from datetime import date, datetime
from typing import Any

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical import Chemical
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer import Customer
from slim_domain.domain.customer.customer_service import CustomerService
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.tr.enums import ProcessingTime, RequestType
from slim_domain.domain.tr.tr_sample_form_parser import (
    SampleRangeBounds,
    TRSampleFormParser,
    _parse_element_symbols,
)
from slim_domain.domain.tr.tr_submission_form_parser import (
    TRSubmissionFormParser,
    _extract_service_date,
    _get_sample_range,
    _parse_emails,
)
from slim_domain.domain.tr.tr_submission import SENTINEL_DATE

# ---------------------------------------------------------------------------
# Stub resolver — returns seeded fixtures regardless of input
# ---------------------------------------------------------------------------

_STUB_CUSTOMER = Customer(
    id=1,
    name="Customer ea6f9",
    street_address="123 Main St",
    city="Springfield",
    state="IL",
    postal_code="62701",
    country="USA",
)
_STUB_CHEMICAL = Chemical(
    id=7,
    name="Hydrochloric Acid",
    metals_prep="Dilute and Shoot",
    silicon_prep="N/A",
    ions_prep="Dilute and Shoot",
)


class _StubResolver:
    def resolve_customer(self, name: str, address: str, address2: str) -> Customer:
        return _STUB_CUSTOMER

    def resolve_chemical(self, form_name: str) -> Chemical:
        return _STUB_CHEMICAL


# ---------------------------------------------------------------------------
# Minimal service fixtures (using in-memory SQLite)
# ---------------------------------------------------------------------------


@pytest.fixture
def session() -> Session:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        # One element so AddElement tests can resolve a symbol
        s.add(_ElementRow(id=14, element_symbol="Si", element_name="Silicon"))
        s.commit()
        yield s


@pytest.fixture
def analysis_svc(session: Session) -> AnalysisService:
    return AnalysisService(session)


@pytest.fixture
def chemical_svc(session: Session) -> ChemicalService:
    return ChemicalService(session)


@pytest.fixture
def element_svc(session: Session) -> ElementService:
    return ElementService(session)


@pytest.fixture
def sample_parser(
    chemical_svc: ChemicalService,
    analysis_svc: AnalysisService,
    element_svc: ElementService,
) -> TRSampleFormParser:
    return TRSampleFormParser(
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=_StubResolver(),
    )


@pytest.fixture
def submission_parser(
    chemical_svc: ChemicalService,
    analysis_svc: AnalysisService,
    element_svc: ElementService,
) -> TRSubmissionFormParser:
    return TRSubmissionFormParser(
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=_StubResolver(),
    )


# ---------------------------------------------------------------------------
# Helper: build a minimal Chemical worksheet
# ---------------------------------------------------------------------------

def _make_chemical_ws() -> Any:
    wb = openpyxl.Workbook()
    ws = wb.active
    # Worksheet tab titles can't contain brackets; the service date comes
    # from the workbook *filename* which is passed separately to the parser.

    ws["D3"] = "Chemical Testing Request Form"
    ws["C10"] = "Customer ea6f9"
    ws["C11"] = "123 Main St"
    ws["C12"] = ""
    ws["I4"] = date(2026, 1, 15)
    ws["P5"] = date(2026, 1, 16)
    ws["P7"] = "JT"
    ws["C13"] = "John Smith"
    ws["B14"] = "Phone :"
    ws["C14"] = "555-1234"
    ws["P49"] = "Rev T"
    ws["H10"] = "PO-12345"
    ws["I11"] = "Visa 1234"
    ws["C15"] = "results@example.com"
    ws["C17"] = ""
    ws["N10"] = "invoice@example.com"
    ws["N11"] = ""

    # Header row = 20; put data in cols 3-15 to set last_col = 14 (B20 = "Sample ID" set after)
    for col in range(3, 16):
        ws.cell(row=20, column=col).value = f"H{col}"
    # Sample range: B20 = "Sample ID" → first_row = 22 (must be set after the loop above)
    ws["B20"] = "Sample ID"
    # One sample in row 22
    ws["B22"] = "S001"
    ws["C22"] = "Hydrochloric Acid"   # chemical name (col 3)
    ws["E22"] = "Si"                  # additional element (col 5)
    ws["K22"] = "Next Day"            # processing time (col 11)
    ws["M22"] = "2:00 PM"             # requested time (col 13)
    ws["N22"] = "Handle carefully"    # notes (col 14)
    # Footer: last row in col B that has data is row 33 (→ last_row = 33 - 11 = 22)
    ws["B33"] = "Total"

    return ws


def _make_wafer_ws() -> Any:
    """Minimal Wafer-shaped worksheet -- real column map confirmed against
    the real Wafer intake form: col3=Reporting Unit, col4=Wafer Size,
    col5=Trace Elements Analysis (# of Elements), col6=Additional Elements,
    col9=Processing Time, col10=Notes. B20's real literal text is
    "Sample ID / Slot Number" (NOT "Sample ID" exactly, unlike Chemical/
    Water) -- this is the real form shape that exposed the
    _get_sample_range prefix-match fix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D3"] = "Wafer Testing Request Form"
    for col in range(3, 12):
        ws.cell(row=20, column=col).value = f"H{col}"
    ws["B20"] = "Sample ID / Slot Number"
    ws["B22"] = "Slot-1"
    ws["C22"] = "atoms/cm^2"       # Reporting Unit (col 3)
    ws["D22"] = "150mm"            # Wafer Size (col 4)
    ws["I22"] = "Next Day"         # processing time (col 9)
    ws["J22"] = "Handle carefully"  # notes (col 10)
    ws["B33"] = "Total"            # footer marker -> last_row = 33 - 11 = 22
    return ws


# ---------------------------------------------------------------------------
# _parse_element_symbols
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("Si,Fe,Cu", ["Si", "Fe", "Cu"]),
        ("Si;Fe;Cu", ["Si", "Fe", "Cu"]),
        ("Si Fe Cu", ["Si", "Fe", "Cu"]),
        ("Si.Fe.Cu", ["Si", "Fe", "Cu"]),
        ("Si, Fe , Cu", ["Si", "Fe", "Cu"]),
        ("", []),
        ("  ", []),
    ],
)
def test_parse_element_symbols(raw: str, expected: list[str]) -> None:
    assert _parse_element_symbols(raw) == expected


# ---------------------------------------------------------------------------
# _get_sample_range
# ---------------------------------------------------------------------------


def test_get_sample_range_exact_sample_id_header_chemical() -> None:
    """Chemical/Water's real B20 is literally "Sample ID" -- exact match,
    first_row = 22."""
    ws = _make_chemical_ws()
    bounds = _get_sample_range(ws)
    assert bounds is not None
    assert bounds.first_row == 22


def test_get_sample_range_prefix_matches_wafer_header_text() -> None:
    """Regression test for a real bug: Wafer's B20 is "Sample ID / Slot
    Number", not "Sample ID" exactly. An exact-equality check silently fell
    through to the row-21 fallback for every real Wafer form, which is the
    sub-header row, not the true header row -- throwing off last_col
    (computed from first_row - 2) and truncating the analysis-column scan.
    Confirmed directly against the real Wafer intake form."""
    ws = _make_wafer_ws()
    bounds = _get_sample_range(ws)
    assert bounds is not None
    assert bounds.first_row == 22
    assert bounds.last_col == 10  # from the real header row (20), not row 19


# ---------------------------------------------------------------------------
# _extract_service_date
# ---------------------------------------------------------------------------


def test_extract_service_date_iso_brackets() -> None:
    assert _extract_service_date("Lab [2026-01-20].xlsx") == date(2026, 1, 20)


def test_extract_service_date_underscore_separator() -> None:
    assert _extract_service_date("Lab [2026_01_20].xlsx") == date(2026, 1, 20)


def test_extract_service_date_no_brackets_returns_sentinel() -> None:
    assert _extract_service_date("Lab Results.xlsx") == SENTINEL_DATE


# ---------------------------------------------------------------------------
# _parse_emails
# ---------------------------------------------------------------------------


def test_parse_emails_single() -> None:
    assert _parse_emails("a@b.com") == ["a@b.com"]


def test_parse_emails_semicolon_separated() -> None:
    assert _parse_emails("a@b.com;c@d.com") == ["a@b.com", "c@d.com"]


def test_parse_emails_filters_invalid() -> None:
    assert _parse_emails("notanemail,a@b.com") == ["a@b.com"]


def test_parse_emails_empty_string() -> None:
    assert _parse_emails("") == []


# ---------------------------------------------------------------------------
# TRSampleFormParser
# ---------------------------------------------------------------------------


def test_sample_parser_builds_sample(
    sample_parser: TRSampleFormParser,
    session: Session,
) -> None:
    ws = _make_chemical_ws()
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.CHEMICAL)

    assert sample.sample_name == "S001"
    assert sample.form_chemical_name == "Hydrochloric Acid"
    assert sample.chemical_id == _STUB_CHEMICAL.id
    assert sample.processing_time == ProcessingTime.NEXT_DAY
    assert sample.requested_time == "2:00 PM"
    assert sample.additional_notes == "Handle carefully"
    assert sample.reporting_units == ""  # not applicable to Chemical samples


def test_sample_parser_resolves_additional_elements(
    sample_parser: TRSampleFormParser,
) -> None:
    ws = _make_chemical_ws()
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.CHEMICAL)
    # "Si" in col 5 should resolve to element_id=14
    assert 14 in sample.additional_element_ids


def test_sample_parser_additional_elements_only_sample(
    sample_parser: TRSampleFormParser,
) -> None:
    """Mocked-up "additional elements only" sample -- confirmed real
    shape (a real Water customer's sample requested nothing but free-text
    additional elements): the Chemical Matrix column still identifies the
    sample (col 3), but the "# of Elements" catalog-panel column (col 4)
    is genuinely blank -- only "Additional Elements (specify)" (col 5)
    is populated. Additional Elements always lives in its own column on
    the real form; this test doesn't rely on any particular real
    submission's data-entry quirks."""
    ws = _make_chemical_ws()
    ws["D22"] = ""  # no catalog "# of Elements" panel selected
    assert ws["E22"].value == "Si"  # additional elements column, from the base fixture

    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.CHEMICAL)

    assert sample.form_chemical_name == "Hydrochloric Acid"  # chemical matrix still identifies the sample
    assert 14 in sample.additional_element_ids


def test_sample_parser_analysis_ids_empty_when_form_analyses_stubbed(
    sample_parser: TRSampleFormParser,
) -> None:
    ws = _make_chemical_ws()
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.CHEMICAL)
    assert sample.analysis_ids == ()


def test_sample_parser_water_uses_constant_chemical_name(
    sample_parser: TRSampleFormParser,
) -> None:
    ws = _make_chemical_ws()
    ws["D3"] = "Water Testing Request Form"
    ws["N22"] = "Next Day"   # WATER reads processing time from col 14
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WATER)
    assert sample.form_chemical_name == "Water"


def test_sample_parser_water_reads_water_package_column(
    sample_parser: TRSampleFormParser,
) -> None:
    """Confirmed real (2026-09-28, a real Water dummy TR run through the
    legacy macro): "Water testing package" (e.g. "Premium package") is a
    real Water-form-only column (3) that bundles a fixed set of panels --
    not a real catalog analysis, so it gets its own TRSample field rather
    than being forced through analysis_ids."""
    ws = _make_chemical_ws()
    ws["D3"] = "Water Testing Request Form"
    ws["N22"] = "Next Day"  # WATER reads processing time from col 14
    ws["C22"] = "Premium package"
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WATER)
    assert sample.water_package == "Premium package"


def test_sample_parser_water_package_empty_when_not_selected(
    sample_parser: TRSampleFormParser,
) -> None:
    ws = _make_chemical_ws()
    ws["D3"] = "Water Testing Request Form"
    ws["N22"] = "Next Day"
    ws["C22"] = None  # the shared fixture defaults this to a chemical name; clear it for Water
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WATER)
    assert sample.water_package == ""


def test_sample_parser_chemical_never_reads_water_package_column(
    sample_parser: TRSampleFormParser,
) -> None:
    """Chemical's own column 3 is the real chemical name, not a package
    selection -- confirms this field stays empty rather than accidentally
    picking up unrelated column-3 content."""
    ws = _make_chemical_ws()
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.CHEMICAL)
    assert sample.water_package == ""


def test_sample_parser_wafer_never_reads_water_package_column(
    sample_parser: TRSampleFormParser,
) -> None:
    """Wafer's own column 3 is reporting_units, a different real field
    that happens to share the index -- confirms no cross-talk."""
    ws = _make_chemical_ws()
    ws["D3"] = "Wafer"
    ws["I22"] = "Next Day"  # WAFER reads processing time from col 9
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WAFER)
    assert sample.water_package == ""


def test_sample_parser_water_resolves_analysis_from_notes_column(
    session: Session,
    chemical_svc: ChemicalService,
    element_svc: ElementService,
) -> None:
    """Confirmed real (2026-09-23, direct user correction): the legacy
    macro scanned "Additional Requests or Notes" for a form-name match
    exactly like any other column -- e.g. a real Folder b65a8 Water form
    requested "Acetate & Formate" only in this free-text column, with no
    structured dropdown selection for it anywhere. Water's notes column
    (17) sits past this fixture's last_col (14, "Processing Time"), so
    this only passes if the notes column is scanned as an extra column,
    not just folded into the existing range.

    Seeds analyses/form_analyses and builds the parser directly, rather
    than using the shared sample_parser/analysis_svc fixtures -- those
    build their AnalysisService (and its startup-loaded cache) before a
    test body runs, too early to see rows seeded inside the test."""
    session.add(_AnalysisRow(id=53, analysis_name="Acetate and Formate", analysis_description=None))
    session.add(_FormAnalysisRow(form_name="Acetate & Formate", analysis_id=53))
    session.commit()
    sample_parser = TRSampleFormParser(
        chemical_service=chemical_svc,
        analysis_service=AnalysisService(session),
        element_service=element_svc,
        input_resolver=_StubResolver(),
    )

    ws = _make_chemical_ws()
    ws["D3"] = "Water Testing Request Form"
    ws["N22"] = "Next Day"           # WATER reads processing time from col 14
    ws["Q22"] = "Acetate & Formate"  # WATER's notes column (17), past last_col
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=14)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WATER)

    assert 53 in sample.analysis_ids


def test_sample_parser_wafer_reads_reporting_units_and_wafer_size(
    sample_parser: TRSampleFormParser,
) -> None:
    """Confirms the real gap-fix: Wafer's "Reporting Unit" column (3) is
    now captured, and "Wafer Size" (column 4) still lands in
    form_chemical_name as before (misleadingly named for this request
    type, but a pre-existing, unchanged mapping)."""
    ws = _make_wafer_ws()
    bounds = SampleRangeBounds(first_row=22, last_row=22, first_col=2, last_col=10)
    sample = sample_parser.build_from_row(ws, row=22, sample_range=bounds, request_type=RequestType.WAFER)

    assert sample.reporting_units == "atoms/cm^2"
    assert sample.form_chemical_name == "150mm"


# ---------------------------------------------------------------------------
# TRSubmissionFormParser
# ---------------------------------------------------------------------------


_CHEMICAL_FILENAME = "Lab Results [2026-01-20].xlsx"


def test_submission_parser_builds_submission(
    submission_parser: TRSubmissionFormParser,
) -> None:
    ws = _make_chemical_ws()
    fixed_dt = datetime(2026, 1, 16, 9, 0)
    sub = submission_parser.build_from_worksheet(ws, download_date=fixed_dt, filename=_CHEMICAL_FILENAME)

    assert sub.customer_id == _STUB_CUSTOMER.id
    assert sub.request_type == RequestType.CHEMICAL
    assert sub.date_submitted == date(2026, 1, 15)
    assert sub.date_received == date(2026, 1, 16)
    assert sub.customer_contact == "John Smith"
    assert sub.customer_phone == "555-1234"
    assert sub.po_information == "PO-12345"
    assert sub.download_date == fixed_dt
    assert sub.technician_initials == "JT"


def test_submission_parser_attaches_samples(
    submission_parser: TRSubmissionFormParser,
) -> None:
    ws = _make_chemical_ws()
    sub = submission_parser.build_from_worksheet(ws, download_date=datetime.now(), filename=_CHEMICAL_FILENAME)
    assert len(sub.samples) == 1
    assert sub.samples[0].sample_name == "S001"


def test_submission_parser_reads_results_email(
    submission_parser: TRSubmissionFormParser,
) -> None:
    ws = _make_chemical_ws()
    sub = submission_parser.build_from_worksheet(ws, download_date=datetime.now(), filename=_CHEMICAL_FILENAME)
    assert "results@example.com" in sub.results_email_main


def test_submission_parser_service_date_from_filename(
    submission_parser: TRSubmissionFormParser,
) -> None:
    ws = _make_chemical_ws()
    sub = submission_parser.build_from_worksheet(ws, download_date=datetime.now(), filename=_CHEMICAL_FILENAME)
    assert sub.service_date == date(2026, 1, 20)


def test_submission_parser_form_customer_fields(
    submission_parser: TRSubmissionFormParser,
) -> None:
    ws = _make_chemical_ws()
    sub = submission_parser.build_from_worksheet(ws, download_date=datetime.now(), filename=_CHEMICAL_FILENAME)
    assert sub.form_customer_name == "Customer ea6f9"
    assert sub.form_customer_address == "123 Main St"
    assert sub.form_customer_address_2 == ""
